from docx import Document

from learner import learner


class Recorder(object):
    def __init__(self):
        self.sentences = []

    def insertCollection(self, sentence):
        self.sentences.append(sentence)


def test_learner_formats_tables_and_paragraphs(tmp_path):
    path = tmp_path / "sheet.docx"
    document = Document()
    document.add_paragraph(":hidden note")
    document.add_paragraph("plain sentence")
    document.add_paragraph("term ** definition")
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "Office"
    table.cell(0, 1).text = "CMPE273"
    table.cell(0, 2).text = "ENG 281"
    document.save(path)

    recorder = Recorder()
    learner(document_path=str(path), inserter=recorder).read()
    assert recorder.sentences[0] == "Office for CMPE273 is ENG 281"
    assert "hidden note" in recorder.sentences
    assert "plain sentence" in recorder.sentences
    assert "term is definition" in recorder.sentences
